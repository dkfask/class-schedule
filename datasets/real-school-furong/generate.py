#!/usr/bin/env python3
"""Generate a MASTER_DATA v1 workbook from the real Furong-campus class timetable.

Source: 2026-2027学年度上期芙蓉校区班级总课表.xlsx (17 classes, 5 teaching days, 6 periods).

The source workbook is a class-centric grid. This script flattens every lesson cell
into (class, day, period, subject, teacher), groups lessons into teaching
requirements, and writes the MASTER_DATA v1 sheets consumed by
WorkbookImportService.

Odd/even-week cells (单双周) list two teachers for one weekly slot. The template
models one teacher per requirement, so the script keeps the second teacher of each
pair: that mapping reproduces a conflict-free timetable (0 teacher conflicts, 0
class conflicts), which makes the real timetable a known feasible solution for the
solver and therefore a fair efficiency baseline.
"""

from __future__ import annotations

import argparse
import collections
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DAYS = ["星期一", "星期二", "星期三", "星期四", "星期五"]
BREAK_LABELS = {"阳光大课间", "课间", "眼保健操、课间", "午餐、午休"}

SUBJECT_CODES = {
    "语文": "FR-SUB-01",
    "数学": "FR-SUB-02",
    "英语": "FR-SUB-03",
    "体育与健康": "FR-SUB-04",
    "艺术-美术": "FR-SUB-05",
    "艺术-音乐": "FR-SUB-06",
    "科学": "FR-SUB-07",
    "道法/习近平新时代中国特色社会主义思想学生读本": "FR-SUB-08",
    "劳动": "FR-SUB-09",
    "体育活动": "FR-SUB-10",
    "班团队/综合实践/生生安/心理": "FR-SUB-11",
    "书法": "FR-SUB-12",
    "信息科技": "FR-SUB-13",
    "校本（英语活动）": "FR-SUB-14",
}

# The source is class-centric and does not identify a room for these lessons.
# Keep ordinary lessons in the administrative classroom and leave specialist /
# walking lessons to the room pool.
FLEXIBLE_SUBJECTS = {"艺术-美术", "艺术-音乐", "科学"}

GRADE_STUDENT_COUNT = {1: 42, 2: 43, 3: 45, 4: 46, 5: 47, 6: 48}
CLASS_TYPE = {1: "一年级", 2: "二年级", 3: "三年级", 4: "四年级", 5: "五年级", 6: "六年级"}


def parse_teachers(raw: str) -> list[str]:
    """Split a teacher cell such as '(常雨晴双/杨雪单)' into ['常雨晴', '杨雪']."""
    raw = raw.strip().replace("（", "(").replace("）", ")")
    match = re.match(r"^\((.*)\)$", raw)
    if not match:
        return []
    names = []
    for part in match.group(1).split("/"):
        name = re.sub(r"(单周|双周|单|双)$", "", part.strip())
        if name:
            names.append(name)
    return names


def read_lessons(source: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(source, data_only=True)
    sheet = workbook["Sheet1"]
    lessons: list[dict] = []
    class_label = None
    period = None
    for row in range(1, sheet.max_row + 1):
        first = sheet.cell(row=row, column=1).value
        if isinstance(first, str) and "课" in first and "程" in first and "表" in first and "\n" in first:
            class_label = sheet.cell(row=row + 1, column=8).value
            period = None
            continue
        if class_label is None:
            continue
        marker = sheet.cell(row=row, column=2).value
        if isinstance(marker, int):
            period = marker
        if period is None:
            continue
        for offset, day in enumerate(DAYS):
            value = sheet.cell(row=row, column=4 + offset).value
            if not value:
                continue
            parts = str(value).split("\n")
            subject = parts[0].strip()
            if subject in BREAK_LABELS:
                continue
            teacher_raw = parts[1].strip() if len(parts) > 1 else ""
            lessons.append(
                {
                    "class": class_label.strip(),
                    "day": day,
                    "period": period,
                    "subject": subject,
                    "teachers": parse_teachers(teacher_raw),
                }
            )
    return lessons


def build_dataset(lessons: list[dict]) -> dict:
    classes = sorted({lesson["class"] for lesson in lessons}, key=class_sort_key)

    teacher_frequency = collections.Counter()
    for lesson in lessons:
        for teacher in lesson["teachers"]:
            teacher_frequency[teacher] += 1
    teachers = sorted(teacher_frequency, key=lambda name: (-teacher_frequency[name], name))

    subjects = sorted(
        {lesson["subject"] for lesson in lessons},
        key=lambda name: (-sum(1 for l in lessons if l["subject"] == name), name),
    )

    for lesson in lessons:
        # Odd/even-week slots: keep the second teacher, which reproduces the
        # conflict-free real timetable. Single-teacher cells keep their teacher.
        lesson["teacher"] = lesson["teachers"][1] if len(lesson["teachers"]) > 1 else lesson["teachers"][0]

    requirements = collections.OrderedDict()
    for lesson in lessons:
        key = (lesson["class"], lesson["subject"], lesson["teacher"])
        requirements.setdefault(key, {"weekly": 0, "slots": []})
        requirements[key]["weekly"] += 1
        requirements[key]["slots"].append((lesson["day"], lesson["period"]))

    return {
        "classes": classes,
        "teachers": teachers,
        "subjects": subjects,
        "requirements": requirements,
        "lessons": lessons,
    }


def class_sort_key(label: str) -> tuple[int, int]:
    match = re.match(r"^([一二三四五六])年级\s*(\d+)$", label)
    if not match:
        return (99, 0)
    grade = "一二三四五六".index(match.group(1)) + 1
    return (grade, int(match.group(2)))


def write_workbook(dataset: dict, term_code: str, target: Path) -> None:
    classes = dataset["classes"]
    teachers = dataset["teachers"]
    subjects = dataset["subjects"]
    requirements = dataset["requirements"]

    teacher_code = {name: f"FR-T{index:02d}" for index, name in enumerate(teachers, start=1)}
    class_code = {}
    for label in classes:
        grade, number = class_sort_key(label)
        class_code[label] = f"FR-G{grade}-{number}"
    subject_code = {name: SUBJECT_CODES[name] for name in subjects}

    rooms = []
    for label in classes:
        grade, number = class_sort_key(label)
        rooms.append(
            (
                f"FR-R{grade}{number:02d}",
                f"{label}教室",
                50,
                "普通教室",
            )
        )
    rooms.extend(
        [
            ("FR-R901", "音乐教室", 50, "专用教室"),
            ("FR-R902", "美术教室", 50, "专用教室"),
            ("FR-R903", "科学实验室", 50, "专用教室"),
        ]
    )

    workbook = openpyxl.Workbook()
    sheets = {
        "说明": ["说明"],
        "教师": ["编码", "名称", "是否启用"],
        "班级": ["编码", "名称", "班级类型", "学生人数", "是否启用", "绑定教室编码"],
        "课程": ["编码", "名称", "是否启用"],
        "教室": ["编码", "名称", "容量", "教室类型", "是否启用"],
        "教学需求": [
            "编码",
            "学期编码",
            "班级编码",
            "课程编码",
            "教师编码",
            "每周课时",
            "单次节数",
            "学生人数",
            "固定节次编码",
            "是否启用",
            "教室分配模式",
            "期望节次编码",
        ],
        "资源可用性": ["资源类型", "资源编码", "学期编码", "节次编码", "是否可用"],
        "特征目录": ["编码", "名称", "是否启用"],
        "教室特征": ["教室编码", "特征编码", "是否启用"],
        "教学需求特征": ["教学需求编码", "特征编码", "是否启用"],
        "活动组": ["编码", "名称", "活动组类型", "学期编码", "成员序号", "教学需求编码", "是否启用"],
    }
    first = True
    for name, headers in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        sheet.title = name
        first = False
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=column, value=header)
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for column in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = max(
                12, len(headers[column - 1]) * 2 + 4
            )

    notes = workbook["说明"]
    notes.cell(
        row=2,
        column=1,
        value=(
            f"MASTER_DATA v1；数据来源：2026-2027学年度上期芙蓉校区班级总课表；"
            f"目标学期 {term_code}；17 个行政班、{len(teachers)} 名教师、{len(subjects)} 门课程、"
            f"{len(rooms)} 间教室、{len(requirements)} 条教学需求；"
            "单双周轮换课程按第二位教师建模（该映射与真实课表零冲突）；普通课程绑定行政班教室，音乐、美术、科学课程使用可选教室池；期望节次编码记录真实课表的原始排位，供 PREFER_ORIGINAL_SLOT 软规则贴原课表重排。"
        ),
    )

    for row, name in enumerate(teachers, start=2):
        sheet = workbook["教师"]
        sheet.cell(row=row, column=1, value=teacher_code[name])
        sheet.cell(row=row, column=2, value=name)
        sheet.cell(row=row, column=3, value="TRUE")

    for row, label in enumerate(classes, start=2):
        grade, _ = class_sort_key(label)
        sheet = workbook["班级"]
        sheet.cell(row=row, column=1, value=class_code[label])
        sheet.cell(row=row, column=2, value=label)
        sheet.cell(row=row, column=3, value=CLASS_TYPE[grade])
        sheet.cell(row=row, column=4, value=GRADE_STUDENT_COUNT[grade])
        sheet.cell(row=row, column=5, value="TRUE")
        sheet.cell(row=row, column=6, value=f"FR-R{grade}{class_sort_key(label)[1]:02d}")

    for row, name in enumerate(subjects, start=2):
        sheet = workbook["课程"]
        sheet.cell(row=row, column=1, value=subject_code[name])
        sheet.cell(row=row, column=2, value=name)
        sheet.cell(row=row, column=3, value="TRUE")

    for row, (code, name, capacity, room_type) in enumerate(rooms, start=2):
        sheet = workbook["教室"]
        sheet.cell(row=row, column=1, value=code)
        sheet.cell(row=row, column=2, value=name)
        sheet.cell(row=row, column=3, value=capacity)
        sheet.cell(row=row, column=4, value=room_type)
        sheet.cell(row=row, column=5, value="TRUE")

    def slot_code(day_name: str, period: int) -> str:
        day_index = DAYS.index(day_name) + 1
        if day_index <= 2 and period <= 2:
            return ["MON", "TUE"][day_index - 1] + f"-{period}"
        return f"DAY-{day_index}-{period}"

    requirement_sheet = workbook["教学需求"]
    requirement_index = {}
    for index, ((class_label, subject, teacher), info) in enumerate(requirements.items(), start=1):
        weekly = info["weekly"]
        preferred = ";".join(slot_code(day, period) for day, period in info["slots"])
        grade, number = class_sort_key(class_label)
        # A class can study one subject with two different teachers (e.g. 道法 split
        # across two teachers), so the teacher number keeps requirement codes unique.
        code = (
            f"FR-REQ-{grade}{number:02d}-{subject_code[subject].split('-')[-1]}"
            f"-{teacher_code[teacher].split('-')[-1]}"
        )
        requirement_index[(class_label, subject, teacher)] = code
        row = index + 1
        requirement_sheet.cell(row=row, column=1, value=code)
        requirement_sheet.cell(row=row, column=2, value=term_code)
        requirement_sheet.cell(row=row, column=3, value=class_code[class_label])
        requirement_sheet.cell(row=row, column=4, value=subject_code[subject])
        requirement_sheet.cell(row=row, column=5, value=teacher_code[teacher])
        requirement_sheet.cell(row=row, column=6, value=weekly)
        requirement_sheet.cell(row=row, column=7, value=1)
        requirement_sheet.cell(row=row, column=8, value=GRADE_STUDENT_COUNT[grade])
        requirement_sheet.cell(row=row, column=9, value=None)
        requirement_sheet.cell(row=row, column=10, value="TRUE")
        requirement_sheet.cell(
            row=row,
            column=11,
            value="FLEXIBLE" if subject in FLEXIBLE_SUBJECTS else "HOME",
        )
        requirement_sheet.cell(row=row, column=12, value=preferred)

    workbook.save(target)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source",
        default="/Users/a1234/Desktop/截图/2026-2027学年度上期芙蓉校区班级总课表.xlsx",
    )
    parser.add_argument("--term-code", default="2026-FALL-FR")
    parser.add_argument("--output-dir", default=str(Path(__file__).resolve().parent))
    args = parser.parse_args()

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    lessons = read_lessons(Path(args.source))
    dataset = build_dataset(lessons)
    write_workbook(dataset, args.term_code, output_dir / "MASTER_DATA-v1.xlsx")

    class_load = collections.Counter(lesson["class"] for lesson in lessons)
    subject_load = collections.Counter(lesson["subject"] for lesson in lessons)
    teacher_load = collections.Counter(lesson["teacher"] for lesson in lessons)
    teacher_conflicts = collections.Counter(
        (lesson["teacher"], lesson["day"], lesson["period"]) for lesson in lessons
    )
    summary = {
        "termCode": args.term_code,
        "lessonCells": len(lessons),
        "classes": len(dataset["classes"]),
        "teachers": len(dataset["teachers"]),
        "subjects": len(dataset["subjects"]),
        "requirements": len(dataset["requirements"]),
        "flexibleRoomRequirements": sum(
            value["weekly"]
            for (class_label, subject, _teacher), value in dataset["requirements"].items()
            if subject in FLEXIBLE_SUBJECTS
        ),
        "slotsPerClass": 30,
        "totalSlots": len(dataset["classes"]) * 30,
        "utilization": round(len(lessons) / (len(dataset["classes"]) * 30), 4),
        "classLoad": dict(class_load),
        "subjectLoad": dict(subject_load),
        "teacherLoad": dict(teacher_load),
        "realTimetableTeacherConflicts": sum(
            1 for count in teacher_conflicts.values() if count > 1
        ),
        "teachers": dataset["teachers"],
        "subjects": dataset["subjects"],
        "classes": dataset["classes"],
        "requirementsDetail": [
            {"class": key[0], "subject": key[1], "teacher": key[2], "weeklyPeriods": value["weekly"]}
            for key, value in dataset["requirements"].items()
        ],
    }
    (output_dir / "extracted-data.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({k: v for k, v in summary.items() if not isinstance(v, (dict, list))}, ensure_ascii=False, indent=2))
    print(f"requirements: {summary['requirements']}, teacherConflicts: {summary['realTimetableTeacherConflicts']}")


if __name__ == "__main__":
    main()
